# -*- coding: utf-8 -*-
# ============================================================
# EXPENSE DIARY
# Полный учебный пайплайн:
# данные -> train/test -> дообучение GPT-2 -> accuracy
# -> проверка слова -> классификация -> CSV -> отчёт
# ============================================================

from pathlib import Path
# Подключаем работу с датой и временем.
from datetime import datetime
# Подключаем поиск похожих слов.
import difflib
# Подключаем работу с таблицами CSV.
import pandas as pd
# Подключаем инструменты для нейросети.
import torch
# Подключаем создание графического интерфейса.
import tkinter as tk
# Подключаем создание Excel-файлов.
from openpyxl import Workbook
# Подключаем оформление таблиц Excel.
from openpyxl.styles import Font, Alignment, Border, Side
# Подключаем подготовку данных для обучения.
from torch.utils.data import Dataset, DataLoader
# Подключаем инструменты для GPT-2.
from transformers import GPT2Tokenizer, GPT2ForSequenceClassification, GPT2Config
# Подключаем оптимизатор для обучения.
from torch.optim import AdamW


# -------------------- ФАЙЛЫ --------------------

BASE_DIR = Path(__file__).resolve().parent

GPT2_DIR = BASE_DIR / "gpt2_local"
MODEL_DIR = BASE_DIR / "expense_model"

TRAIN_FILE = BASE_DIR / "training_data.csv"
EXPENSES_FILE = BASE_DIR / "monthly_expenses.csv"

WORDS_FILE = BASE_DIR / "dictionaries" / "english_words.txt"
PRODUCTS_FILE = BASE_DIR / "dictionaries" / "product_names.txt"


# -------------------- КАТЕГОРИИ --------------------

categories = [
    "GROCERIES",
    "TRANSPORT",
    "CLOTHING",
    "BEAUTY",
    "HOME",
    "ELECTRONICS",
    "HEALTH",
    "ENTERTAINMENT",
    "RESTAURANTS",
    "OTHER"
]

label2id = {
    name: number
    for number, name in enumerate(categories)
}

id2label = {
    number: name
    for number, name in enumerate(categories)
}
# Для проекта используем CPU.
device = torch.device("cpu")

# -------------------- ПОДГОТОВКА ДАННЫХ --------------------

# Читает и подготавливает данные для обучения и тестирования.
def prepare_data():

    if not TRAIN_FILE.exists():
        raise FileNotFoundError(
            "training_data.csv was not found."
        )

    data = pd.read_csv(TRAIN_FILE)

    # Проверяем структуру датасета.
    if not {"text", "category"}.issubset(data.columns):
        raise ValueError(
            "training_data.csv must contain "
            "'text' and 'category'."
        )

    data = data.dropna(
        subset=["text", "category"]
    ).copy()

    data["text"] = (
        data["text"]
        .astype(str)
        .str.strip()
    )

    data["category"] = (
        data["category"]
        .astype(str)
        .str.strip()
        .str.upper()
    )

    # Категорию превращаем в число.
    data["label"] = data["category"].map(label2id)

    if data["label"].isna().any():
        raise ValueError(
            "Unknown category in training_data.csv."
        )

    data["label"] = data["label"].astype(int)

    train_parts = []
    test_parts = []

    # Каждую категорию отдельно делим 80% / 20%.
    for category in categories:

        part = data[
            data["category"] == category
        ].sample(
            frac=1,
            random_state=42
        ).reset_index(drop=True)

        split = int(len(part) * 0.8)

        train_parts.append(
            part.iloc[:split]
        )

        test_parts.append(
            part.iloc[split:]
        )

    train_data = pd.concat(
        train_parts,
        ignore_index=True
    )

    test_data = pd.concat(
        test_parts,
        ignore_index=True
    )

    print(
        f"Train examples: {len(train_data)}"
    )

    print(
        f"Test examples: {len(test_data)}"
    )

    return train_data, test_data


# -------------------- DATASET --------------------

class ExpenseDataset(Dataset):

    # Сохраняет тексты, категории и токенизатор.
    def __init__(self, data, tokenizer):

        self.texts = data["text"].tolist()
        self.labels = data["label"].tolist()
        self.tokenizer = tokenizer

    # Возвращает количество примеров.
    def __len__(self):

        return len(self.texts)

    # Преобразует один пример в данные для модели.
    def __getitem__(self, index):

        encoded = self.tokenizer(
            self.texts[index],
            truncation=True,
            padding="max_length",
            max_length=32,
            return_tensors="pt"
        )

        return {
            "input_ids":
                encoded["input_ids"].squeeze(0),

            "attention_mask":
                encoded["attention_mask"].squeeze(0),

            "labels":
                torch.tensor(
                    self.labels[index],
                    dtype=torch.long
                )
        }


# -------------------- ОЦЕНКА МОДЕЛИ --------------------

# Проверяет точность модели на тестовых данных.
def evaluate_model(model, loader):

    model.eval()

    correct = 0
    total = 0

    with torch.no_grad():

        for batch in loader:

            output = model(
                input_ids=
                    batch["input_ids"].to(device),

                attention_mask=
                    batch["attention_mask"].to(device)
            )

            predictions = torch.argmax(
                output.logits,
                dim=1
            )

            labels = batch["labels"].to(device)

            correct += (
                predictions == labels
            ).sum().item()

            total += labels.size(0)

    return correct / total


# -------------------- ДООБУЧЕНИЕ GPT-2 --------------------

# Дообучает GPT-2 и сохраняет результат.
def train_model():

    if not (GPT2_DIR / "config.json").exists():

        raise FileNotFoundError(
            "Folder gpt2_local was not found."
        )

    train_data, test_data = prepare_data()

    # Загружаем локальный токенизатор GPT-2.
    tokenizer = GPT2Tokenizer.from_pretrained(
        GPT2_DIR,
        local_files_only=True
    )

    tokenizer.pad_token = tokenizer.eos_token

    train_dataset = ExpenseDataset(
        train_data,
        tokenizer
    )

    test_dataset = ExpenseDataset(
        test_data,
        tokenizer
    )

    train_loader = DataLoader(
        train_dataset,
        batch_size=4,
        shuffle=True
    )

    test_loader = DataLoader(
        test_dataset,
        batch_size=4,
        shuffle=False
    )

    # Загружаем конфигурацию базовой GPT-2.
    config = GPT2Config.from_pretrained(
        GPT2_DIR,
        local_files_only=True
    )

    config.num_labels = len(categories)
    config.label2id = label2id
    config.id2label = id2label
    config.pad_token_id = tokenizer.pad_token_id

    # Создаём GPT-2 классификатор.
    model = GPT2ForSequenceClassification.from_pretrained(
        GPT2_DIR,
        config=config,
        local_files_only=True
    ).to(device)

    optimizer = AdamW(
        model.parameters(),
        lr=5e-5
    )

    epochs = 5

    # Дообучение.
    for epoch in range(epochs):

        model.train()
        total_loss = 0

        for batch in train_loader:

            optimizer.zero_grad()

            output = model(
                input_ids=
                    batch["input_ids"].to(device),

                attention_mask=
                    batch["attention_mask"].to(device),

                labels=
                    batch["labels"].to(device)
            )

            loss = output.loss

            loss.backward()
            optimizer.step()

            total_loss += loss.item()

        average_loss = (
            total_loss / len(train_loader)
        )

        print(
            f"Epoch {epoch + 1}/{epochs} "
            f"- Loss: {average_loss:.4f}"
        )

    # Проверяем качество на test.
    accuracy = evaluate_model(
        model,
        test_loader
    )

    print(
        f"Test Accuracy: {accuracy:.2%}"
    )

    # Сохраняем дообученную модель.
    MODEL_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    model.save_pretrained(MODEL_DIR)
    tokenizer.save_pretrained(MODEL_DIR)

    print(
        "Fine-tuned model saved "
        "in expense_model."
    )

    return tokenizer, model


# -------------------- ЗАГРУЗКА МОДЕЛИ --------------------

# Загружает ранее обученную модель.
def load_model():

    tokenizer = GPT2Tokenizer.from_pretrained(
        MODEL_DIR,
        local_files_only=True
    )

    if tokenizer.pad_token is None:
        tokenizer.pad_token = tokenizer.eos_token

    model = GPT2ForSequenceClassification.from_pretrained(
        MODEL_DIR,
        local_files_only=True
    ).to(device)

    model.config.pad_token_id = (
        tokenizer.pad_token_id
    )

    return tokenizer, model


# Если модель уже обучена, повторно её не обучаем.
if (MODEL_DIR / "config.json").exists():

    print("Loading fine-tuned model...")

    tokenizer, model = load_model()

else:

    print("First launch: fine-tuning GPT-2...")

    tokenizer, model = train_model()


# -------------------- ЛОКАЛЬНЫЕ СЛОВАРИ --------------------

# Загружает слова из локального файла.
def load_dictionary(path):

    if not path.exists():

        raise FileNotFoundError(
            f"Dictionary was not found: {path}"
        )

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:

        return {
            line.strip().lower()
            for line in file
            if line.strip()
        }


# В этой версии словари обязательны.
english_words = load_dictionary(
    WORDS_FILE
)

product_names = load_dictionary(
    PRODUCTS_FILE
)


# -------------------- ПРОВЕРКА СЛОВА --------------------

# Проверяет, есть ли товар в словарях.
def is_known_product(product):

    text = product.lower().strip()

    # Если полное название есть в списке товаров.
    if text in product_names:
        return True

    # Проверяем отдельные слова.
    parts = text.replace("-", " ").split()

    if (
        parts
        and all(
            word in english_words
            for word in parts
        )
    ):
        return True

    return False

# Ищет похожее правильное написание.
def find_correction(product):

    text = product.lower().strip()

    # Сначала ищем похожее название товара.
    match = difflib.get_close_matches(
        text,
        product_names,
        n=1,
        cutoff=0.80
    )

    if match:
        return match[0]

    # Если введено одно слово,
    # ищем похожее слово в английском словаре.
    parts = text.replace("-", " ").split()

    if len(parts) == 1:

        match = difflib.get_close_matches(
            text,
            english_words,
            n=1,
            cutoff=0.82
        )

        if match:
            return match[0]

    return None

# Проверяет название и предлагает исправление.
def check_product(product):
    text = product.lower().strip()
    # Числа например "54" не являются товаром.
    if not any(
        letter.isalpha()
        for letter in text
    ):
        return None

    # Если слово уже правильное.
    if is_known_product(text):
        return text

    # Ищем возможное исправление.
    correction = find_correction(text)

    if correction:

        answer = input(
            f"Did you mean "
            f"'{correction}'? (y/n): "
        ).strip().lower()

        if answer == "y":
            # Исправленное слово тоже проверяем.
            if is_known_product(correction):
                return correction
    # Если слова нет в словарях —
    # оно НЕ передаётся GPT-2.
    return None

# -------------------- КЛАССИФИКАЦИЯ GPT-2 --------------------
# Определяет категорию и показатель Confidence.
def classify_product(product):
    inputs = tokenizer(
        product,
        return_tensors="pt",
        truncation=True,
        padding=True,
        max_length=32
    )
    inputs = {
        key: value.to(device)
        for key, value in inputs.items()
    }
    model.eval()
    with torch.no_grad():
        output = model(**inputs)
    probabilities = torch.softmax(
        output.logits,
        dim=1
    )
    confidence, prediction = torch.max(
        probabilities,
        dim=1
    )
    category = id2label[
        prediction.item()
    ]
    return category, confidence.item()

# -------------------- СОХРАНЕНИЕ --------------------

# Сохраняет покупку в CSV-файл.
def save_expense(
    month,
    product,
    price,
    category
):
    row = pd.DataFrame([{
        "Month": month,
        "Product": product,
        "Price": price,
        "Category": category
    }])
    if EXPENSES_FILE.exists():
        row.to_csv(
            EXPENSES_FILE,
            mode="a",
            header=False,
            index=False
        )
    else:
        row.to_csv(
            EXPENSES_FILE,
            index=False
        )

# -------------------- ОТЧЁТ --------------------
# Выбирает расходы за указанный месяц.
def get_month_data(month):
    if not EXPENSES_FILE.exists():
        return None
    data = pd.read_csv(EXPENSES_FILE)
    month_data = data[
        data["Month"].astype(str).str.lower()
        == month.lower()
    ].copy()
    if month_data.empty:
        return None
    month_data["Price"] = pd.to_numeric(
        month_data["Price"],
        errors="coerce"
    ).fillna(0)
    return month_data

# Создаёт и сохраняет отчёт Excel.
def save_month_report_excel(month, month_data):
    # Сохраняем итоговый отчёт Excel из 3 колонок.
    reports_dir = BASE_DIR / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)

    year = datetime.now().year
    report_file = reports_dir / f"{month}_{year}_expenses.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = month

    year = datetime.now().year

    # Заголовок отчёта над таблицей.
    sheet.merge_cells("A1:C1")
    title_cell = sheet["A1"]
    title_cell.value = f"{month.upper()} {year}"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )
    sheet.append(["CATEGORY", "PRODUCT", "PRICE"])
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
    thin = Side(style="thin")
    thick = Side(style="thick")
    current_row = 3
    for category in categories:
        category_data = month_data[
            month_data["Category"] == category
        ]
        if category_data.empty:
            continue
        first_row = current_row

        # Записываем товары одной категории подряд.
        for _, expense in category_data.iterrows():
            sheet.cell(
                row=current_row,
                column=2,
                value=str(expense["Product"])
            )
            price_cell = sheet.cell(
                row=current_row,
                column=3,
                value=float(expense["Price"])
            )
            price_cell.number_format = "0.00"
            current_row += 1
        last_row = current_row - 1
        # Название категории — одно и по центру группы.
        if first_row < last_row:
            sheet.merge_cells(
                start_row=first_row,
                start_column=1,
                end_row=last_row,
                end_column=1
            )
        category_cell = sheet.cell(
            row=first_row,
            column=1,
            value=category
        )
        category_cell.font = Font(bold=True)
        category_cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        # TOTAL слева, сама сумма — в третьей колонке.
        total_row = current_row
        sheet.cell(
            row=total_row,
            column=1,
            value="TOTAL"
        )
        total_price = sheet.cell(
            row=total_row,
            column=3,
            value=float(category_data["Price"].sum())
        )
        total_price.number_format = "0.00"

        for column in range(1, 4):
            cell = sheet.cell(
                row=total_row,
                column=column
            )
            cell.font = Font(bold=True)

            # Над TOTAL тонкая линия, под TOTAL жирная.
            cell.border = Border(
                top=thin,
                bottom=thick
            )
        current_row += 1
    # Общая сумма месяца.
    month_total_row = current_row
    sheet.cell(
        row=month_total_row,
        column=1,
        value="MONTH TOTAL"
    )
    month_total_price = sheet.cell(
        row=month_total_row,
        column=3,
        value=float(month_data["Price"].sum())
    )
    month_total_price.number_format = "0.00"
    for column in range(1, 4):
        cell = sheet.cell(
            row=month_total_row,
            column=column
        )
        cell.font = Font(bold=True)
        cell.border = Border(
            top=thick,
            bottom=thick
        )
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 14

    workbook.save(report_file)
    return report_file

# Показывает месячный отчёт в окне.
def show_total(month):
    # Показываем такой же отчёт в окне tkinter.
    month_data = get_month_data(month)
    if month_data is None:
        print("No expenses for this month.")
        return
    report_file = save_month_report_excel(
        month,
        month_data
    )
    window = tk.Tk()
    window.title(f"Expenses - {month}")
    year = datetime.now().year
    # Месяц и год по центру над всей таблицей.
    tk.Label(
        window,
        text=f"{month.upper()} {year}",
        font=("Arial", 14, "bold"),
        padx=10,
        pady=10
    ).grid(
        row=0,
        column=0,
        columnspan=3,
        sticky="nsew"
    )
    headers = ["CATEGORY", "PRODUCT", "PRICE"]
    for column, title in enumerate(headers):
        tk.Label(
            window,
            text=title,
            font=("Arial", 11, "bold"),
            borderwidth=1,
            relief="solid",
            padx=10,
            pady=6
        ).grid(
            row=1,
            column=column,
            sticky="nsew"
        )
    current_row = 2
    for category in categories:
        category_data = month_data[
            month_data["Category"] == category
        ]
        if category_data.empty:
            continue
        number_of_products = len(category_data)
        # Категория занимает одну большую ячейку
        # и находится по центру списка товаров.
        tk.Label(
            window,
            text=category,
            font=("Arial", 10, "bold"),
            borderwidth=1,
            relief="solid",
            padx=10,
            pady=5
        ).grid(
            row=current_row,
            column=0,
            rowspan=number_of_products,
            sticky="nsew"
        )
        for offset, (_, expense) in enumerate(
            category_data.iterrows()
        ):
            row_number = current_row + offset
            tk.Label(
                window,
                text=str(expense["Product"]),
                borderwidth=1,
                relief="solid",
                padx=10,
                pady=5
            ).grid(
                row=row_number,
                column=1,
                sticky="nsew"
            )

            tk.Label(
                window,
                text=f'{float(expense["Price"]):.2f}',
                borderwidth=1,
                relief="solid",
                padx=10,
                pady=5,
                anchor="e"
            ).grid(
                row=row_number,
                column=2,
                sticky="nsew"
            )

        current_row += number_of_products

        category_total = float(
            category_data["Price"].sum()
        )
        # TOTAL в первой колонке,
        # сумма в третьей колонке.
        for column in range(3):
            value = ""

            if column == 0:
                value = "TOTAL"
            elif column == 2:
                value = f"{category_total:.2f}"

            tk.Label(
                window,
                text=value,
                font=("Arial", 10, "bold"),
                borderwidth=2,
                relief="solid",
                padx=10,
                pady=6,
                anchor="w" if column == 0 else (
                    "e" if column == 2 else "center"
                )
            ).grid(
                row=current_row,
                column=column,
                sticky="nsew"
            )

        current_row += 1

    month_total = float(
        month_data["Price"].sum()
    )

    # Общая сумма месяца:
    # название слева, сумма справа.
    for column in range(3):
        value = ""

        if column == 0:
            value = "MONTH TOTAL"
        elif column == 2:
            value = f"{month_total:.2f}"

        tk.Label(
            window,
            text=value,
            font=("Arial", 11, "bold"),
            borderwidth=3,
            relief="solid",
            padx=10,
            pady=7,
            anchor="w" if column == 0 else (
                "e" if column == 2 else "center"
            )
        ).grid(
            row=current_row,
            column=column,
            sticky="nsew"
        )

    window.columnconfigure(0, weight=2)
    window.columnconfigure(1, weight=3)
    window.columnconfigure(2, weight=1)

    print(f"Report saved: {report_file}")

    window.mainloop()

# -------------------- МЕСЯЦ --------------------

# Запрашивает и проверяет название месяца.
def get_month():
    months = [
        "January",
        "February",
        "March",
        "April",
        "May",
        "June",
        "July",
        "August",
        "September",
        "October",
        "November",
        "December"
    ]
    while True:
        value = input(
            "Enter month "
            "(example: October): "
        ).strip()
        month = value.capitalize()
        if month in months:
            return month
        print(
            "Enter a correct month name, "
            "for example: October."
        )

# -------------------- ОСНОВНАЯ ПРОГРАММА --------------------

print()
print("EXPENSE DIARY")
print("Fine-tuned GPT-2 classifier")
print("Spelling dictionaries: loaded")
print()
month = get_month()
print(
    f"Selected month: {month}"
)
while True:
    product = input(
        "\nProduct / total / month / exit: "
    ).strip()
    # Выход.
    if product.lower() == "exit":

        print("Goodbye!")
        break
    # Отчёт за текущий месяц.
    if product.lower() == "total":

        show_total(month)
        continue
    # Можно перейти в другой месяц и продолжить добавлять расходы.
    if product.lower() == "month":

        month = get_month()
        print(
            f"Selected month: {month}"
        )
        continue
    # Пустая строка.
    if not product:
        print(
            "Enter a product name."
        )
        continue
    # СНАЧАЛА проверяем слово.
    checked_product = check_product(
        product
    )
    # Неизвестное слово останавливаем здесь.
    # Цена не спрашивается.
    # GPT-2 не вызывается.
    if checked_product is None:
        print(
            "Unknown product. "
            "Please check the spelling."
        )
        continue
    product = checked_product
    # Затем спрашиваем цену.
    # При ошибке повторяется только цена.
    while True:
        try:
            price = float(
                input("Price: ")
                .replace(",", ".")
            )
            if price <= 0:
                print(
                    "Price must be "
                    "greater than 0."
                )
                continue
            break
        except ValueError:
            print(
                "Enter a number "
                "for the price."
            )
    # Только проверенный товар
    # отправляется в GPT-2.
    category, confidence = (
        classify_product(product)
    )
    save_expense(
        month,
        product,
        price,
        category
    )
    # После нового ввода товара автоматически
    # обновляем Excel-отчёт выбранного месяца.
    month_data = get_month_data(month)
    if month_data is not None:
        save_month_report_excel(
            month,
            month_data
        )
    print(
        f"Category: {category}"
    )
    print(
        f"Confidence: "
        f"{confidence:.1%}"
    )
    print(
        "Expense saved."
    )
